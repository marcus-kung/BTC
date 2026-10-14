
# ============================================================
# SECTION 1: IMPORT LIBRARIES
# ============================================================
# Import the tools needed to run the program.
#
# argparse: Allows us to provide settings when running the script.
# csv: Reads and writes CSV files used to save progress.
# random: Adds small random delays between Yahoo Finance requests.
# time: Allows the program to pause between requests.
# Path: Helps manage file and folder locations.
# openpyxl: Allows Python to read and edit Excel workbooks.
# yfinance: Retrieves company information from Yahoo Finance.

from __future__ import annotations

import argparse
import csv
import random
import time
from pathlib import Path

from openpyxl import load_workbook
import yfinance as yf


# ============================================================
# SECTION 2: DEFINE CONSTANTS AND HELPER FUNCTIONS
# ============================================================
# Define which company details we want to retrieve.
#
# FIELDS contains the Excel columns that can be updated.
# CACHE_FIELDS defines the information stored in our progress
# file, including whether a lookup succeeded or failed.
#
# clean() ensures that data is consistently stored as text.
# It removes unwanted spaces and converts empty cells into
# empty strings rather than the Python value None.

FIELDS = ('Name', 'Sector', 'Industry')

CACHE_FIELDS = (
    'Ticker',
    'Name',
    'Sector',
    'Industry',
    'Status',
    'Error'
)


def clean(value):
    return str(value).strip() if value is not None else ''


# ============================================================
# SECTION 3: READ AND SAVE PROGRESS
# ============================================================
# Yahoo Finance lookups can take a long time when processing
# thousands of stocks.
#
# To avoid repeating completed work, we store results in a
# separate CSV file called a cache.
#
# read_cache() loads previously saved company information into
# a dictionary, using each stock ticker as its unique key.
#
# write_cache() saves the dictionary back into a CSV file.
# Results are first written to a temporary file, which then
# replaces the existing cache. This helps protect saved
# progress if the program stops unexpectedly.


def read_cache(path):
    if not path.exists():
        return {}

    with path.open(newline='', encoding='utf-8-sig') as f:
        return {
            r['Ticker']: r
            for r in csv.DictReader(f)
            if r.get('Ticker')
        }


def write_cache(path, cache):
    path.parent.mkdir(parents=True, exist_ok=True)

    temp = path.with_suffix('.tmp')

    with temp.open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=CACHE_FIELDS)
        writer.writeheader()

        for ticker in sorted(cache):
            writer.writerow({
                col: cache[ticker].get(col, '')
                for col in CACHE_FIELDS
            })

    temp.replace(path)


# ============================================================
# SECTION 4: RETRIEVE COMPANY INFORMATION FROM YAHOO FINANCE
# ============================================================
# This function takes a stock ticker and requests its company
# profile from Yahoo Finance.
#
# For example, the ticker "AAPL" should return information
# about Apple, including its company name, sector and industry.
#
# Each lookup is assigned one of four possible outcomes:
#
# complete: All three company fields were retrieved.
# partial: Some company information is missing.
# unresolved: Yahoo Finance did not return a valid profile.
# error: An exception occurred during the request.
#
# The try/except structure prevents an individual failed
# request from immediately crashing the entire program.


def query_yahoo(ticker):
    try:
        info = yf.Ticker(ticker).get_info()

        if not isinstance(info, dict):
            return dict(
                Ticker=ticker,
                Name='',
                Sector='',
                Industry='',
                Status='unresolved',
                Error='No profile'
            )

        company = clean(info.get('longName') or info.get('shortName'))
        sector = clean(info.get('sector'))
        industry = clean(info.get('industry'))

        return dict(
            Ticker=ticker,
            Name=company,
            Sector=sector,
            Industry=industry,
            Status='complete' if all(
                (company, sector, industry)
            ) else 'partial',
            Error=''
        )

    except Exception as exc:
        return dict(
            Ticker=ticker,
            Name='',
            Sector='',
            Industry='',
            Status='error',
            Error=f'{type(exc).__name__}: {str(exc)[:160]}'
        )


# ============================================================
# SECTION 5: CONFIGURE PROGRAM SETTINGS AND FILE PATHS
# ============================================================
# The main() function controls the overall program.
#
# We use argparse so that settings can be changed directly
# from the terminal without modifying the Python code.
#
# For example, users can specify:
# - The Excel workbook containing the stock rankings.
# - Where the completed workbook should be saved.
# - How many new stocks to process in one run.
# - How long to wait between Yahoo Finance requests.
# - Whether unsuccessful lookups should be attempted again.
#
# Default output files are placed in an "output" folder.
#
# The program also checks whether an older progress file
# exists and copies it to the new location if necessary.
# This prevents previously completed lookups from being lost.
#
# Finally, the source and output paths must differ so that
# the original Excel workbook is never overwritten.


def main():
    p = argparse.ArgumentParser(
        description='Enrich stock rankings with Yahoo Finance company profiles'
    )

    p.add_argument(
        '--input',
        required=True,
        type=Path,
        help='Source workbook (never modified)'
    )

    p.add_argument(
        '--output',
        type=Path,
        help='Final workbook (default: output/stock_metadata.xlsx next to script)'
    )

    p.add_argument(
        '--cache',
        type=Path,
        help='Progress CSV (default: output/lookup_cache.csv next to script)'
    )

    p.add_argument(
        '--limit',
        type=int,
        default=0,
        help='Max NEW Yahoo lookups this run; 0 means all'
    )

    p.add_argument(
        '--delay',
        type=float,
        default=2.0,
        help='Seconds between requests (default 2)'
    )

    p.add_argument(
        '--batch-size',
        type=int,
        default=50,
        help='Number of requests between longer pauses (default 50)'
    )

    p.add_argument(
        '--batch-pause',
        type=float,
        default=120,
        help='Seconds to pause between batches (default 120)'
    )

    p.add_argument(
        '--rate-limit-pause',
        type=float,
        default=900,
        help='Seconds to wait after HTTP 429 (default 900)'
    )

    p.add_argument(
        '--max-rate-limit-retries',
        type=int,
        default=3,
        help='Retries per ticker after 429 before stopping safely'
    )

    p.add_argument(
        '--retry-failed',
        action='store_true',
        help='Retry earlier partial/error/unresolved results'
    )

    args = p.parse_args()

    output_dir = Path(__file__).resolve().parent / 'output'

    args.output = args.output or output_dir / 'stock_metadata.xlsx'
    args.cache = args.cache or output_dir / 'lookup_cache.csv'

    previous_cache = Path(__file__).resolve().parent / 'yahoo_lookup_progress.csv'

    if not args.cache.exists() and previous_cache.exists():
        import shutil

        args.cache.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(previous_cache, args.cache)

        print(
            f'Migrated existing lookup cache to {args.cache}',
            flush=True
        )

    if args.output.resolve() == args.input.resolve():
        p.error('Output must differ from input to protect the source workbook')


    # ========================================================
    # SECTION 6: LOAD THE EXCEL WORKBOOK AND EXTRACT TICKERS
    # ========================================================
    # Open the input Excel workbook and find the worksheet
    # called "Full Rankings".
    #
    # The worksheet must contain the following column headings:
    # Ticker, Name, Sector and Industry.
    #
    # The program identifies each column's position using its
    # heading rather than assuming a fixed column order.
    #
    # It then extracts all non-empty stock tickers.
    #
    # dict.fromkeys() removes duplicate tickers while keeping
    # their original order, ensuring each company is queried
    # only once during a run.
    #
    # Previously saved results are loaded from the cache.
    # Finally, the program checks that user-supplied settings
    # such as delays and batch sizes are valid.

    print('Loading workbook...')

    workbook = load_workbook(args.input)

    if 'Full Rankings' not in workbook.sheetnames:
        raise ValueError('Workbook must contain a Full Rankings sheet')

    sheet = workbook['Full Rankings']

    header = [clean(cell.value) for cell in sheet[1]]

    cols = {h: i for i, h in enumerate(header)}

    if not all(h in cols for h in ('Ticker', *FIELDS)):
        raise ValueError(f'Missing expected columns: {header}')

    rows = list(sheet.iter_rows(min_row=2, values_only=True))

    tickers = list(dict.fromkeys(
        clean(row[cols['Ticker']])
        for row in rows
        if clean(row[cols['Ticker']])
    ))

    cache = read_cache(args.cache)

    if (
        args.batch_size < 1
        or args.delay < 0
        or args.batch_pause < 0
        or args.rate_limit_pause < 0
        or args.max_rate_limit_retries < 0
        or args.limit < 0
    ):
        p.error(
            'Batch size must be positive; delays, retry count, '
            'and limit must be nonnegative'
        )


    # ========================================================
    # SECTION 7: PROCESS STOCKS AND CONTROL REQUEST FREQUENCY
    # ========================================================
    # Loop through the list of stock tickers and retrieve
    # their company information.
    #
    # Before requesting data, the program checks whether
    # the ticker already exists in the cache.
    #
    # By default, previously attempted tickers are skipped.
    # Using --retry-failed allows incomplete or failed
    # lookups to be attempted again.
    #
    # The program processes stocks in batches to reduce
    # the frequency of requests sent to Yahoo Finance.
    #
    # After each batch, it:
    # 1. Saves the latest progress.
    # 2. Pauses for the configured batch interval.
    # 3. Continues processing the next group of stocks.
    #
    # The --limit setting can also restrict how many new
    # requests are performed in a single run.


    new_count = 0
    batch_count = 0

    try:
        for index, ticker in enumerate(tickers, 1):
            existing = cache.get(ticker)

            if existing and (
                not args.retry_failed
                or existing.get('Status') == 'complete'
            ):
                continue

            if args.limit and new_count >= args.limit:
                break

            if batch_count >= args.batch_size:
                write_cache(args.cache, cache)

                print(
                    f'Batch complete. Sleeping '
                    f'{args.batch_pause:.0f}s before continuing...',
                    flush=True
                )

                time.sleep(args.batch_pause)
                batch_count = 0


            # ====================================================
            # SECTION 8: HANDLE YAHOO FINANCE RATE LIMITS
            # ====================================================
            # Yahoo Finance may temporarily reject requests when
            # too many are sent in a short period.
            #
            # This is commonly reported as HTTP error 429,
            # meaning "Too Many Requests".
            #
            # If a rate-limit error occurs, the program waits
            # before attempting the same ticker again.
            #
            # The waiting period doubles after each unsuccessful
            # retry. This is called exponential backoff.
            #
            # For example, with a starting pause of 900 seconds:
            # First retry:  900 seconds
            # Second retry: 1800 seconds
            # Third retry:  3600 seconds
            #
            # If Yahoo Finance continues to reject requests,
            # the program stops safely and preserves its cache.
            #
            # Successful requests are added to the cache, and
            # progress is saved every 20 completed lookups.


            retries = 0

            while True:
                result = query_yahoo(ticker)

                error = result.get('Error', '').lower()

                rate_limited = (
                    result['Status'] == 'error'
                    and (
                        '429' in error
                        or 'ratelimit' in error
                        or 'rate limit' in error
                        or 'too many requests' in error
                    )
                )

                if not rate_limited:
                    break

                if retries >= args.max_rate_limit_retries:
                    print(
                        f'Yahoo Finance continues to rate-limit '
                        f'after {retries} retries. '
                        'Stopping safely; run again later.',
                        flush=True
                    )
                    break

                retries += 1

                write_cache(args.cache, cache)

                pause = args.rate_limit_pause * (2 ** (retries - 1))

                print(
                    f'Rate limited on {ticker}. Waiting {pause:.0f}s '
                    f'(retry {retries}/{args.max_rate_limit_retries})...',
                    flush=True
                )

                time.sleep(pause)

            if rate_limited:
                break

            cache[ticker] = result

            new_count += 1
            batch_count += 1

            print(
                f'{index}/{len(tickers)} {ticker}: '
                f'{result["Status"]} | {result["Name"][:60]}',
                flush=True
            )

            if new_count % 20 == 0:
                write_cache(args.cache, cache)

                print(
                    f'Checkpoint saved: {len(cache)} symbols',
                    flush=True
                )

            time.sleep(args.delay + random.uniform(0, 0.3))

    except KeyboardInterrupt:
        print('Interrupted. Saving current progress.', flush=True)

    finally:
        write_cache(args.cache, cache)


    # ========================================================
    # SECTION 9: UPDATE MISSING COMPANY DETAILS IN EXCEL
    # ========================================================
    # Now that Yahoo Finance lookups are complete, use the
    # cached information to fill missing Excel cells.
    #
    # For each stock, the program:
    # 1. Reads its ticker from the worksheet.
    # 2. Finds the corresponding company in the cache.
    # 3. Checks the Name, Sector and Industry columns.
    # 4. Fills cells only when they are currently empty.
    #
    # IMPORTANT: Existing company information is preserved.
    # The program does not overwrite classifications or
    # modify the original stock rankings.
    #
    # The variable changed counts how many empty cells
    # have been successfully populated.
    #
    # Finally, the completed workbook is saved to the
    # designated output file.


    changed = 0

    for row_number, row in enumerate(
        sheet.iter_rows(min_row=2),
        start=2
    ):
        ticker = clean(row[cols['Ticker']].value)

        record = cache.get(ticker, {})

        for field in FIELDS:
            cell = row[cols[field]]

            addition = clean(record.get(field, ''))

            if not clean(cell.value) and addition:
                cell.value = addition
                changed += 1

    print('Exporting edited workbook...')

    args.output.parent.mkdir(parents=True, exist_ok=True)

    workbook.save(args.output)


    # ========================================================
    # SECTION 10: EXPORT UNRESOLVED STOCKS AND PRINT SUMMARY
    # ========================================================
    # Create a separate CSV report containing companies that
    # do not have a complete Yahoo Finance profile.
    #
    # This includes:
    # - Stocks with missing company details.
    # - Stocks whose Yahoo Finance requests failed.
    # - Stocks that have not yet been looked up.
    #
    # This report makes it easier to identify which companies
    # may require further investigation or manual corrections.
    #
    # The program finishes by printing:
    # - How many Excel cells were filled.
    # - Where the completed workbook was saved.
    # - Where the progress cache was saved.
    # - Where the unresolved company report was saved.
    #
    # Because progress is cached, the same command can be
    # run again to continue processing remaining stocks.


    failed_path = args.output.parent / 'unresolved.csv'

    with failed_path.open('w', newline='', encoding='utf-8') as f:
        writer = csv.DictWriter(f, fieldnames=CACHE_FIELDS)
        writer.writeheader()

        for ticker in tickers:
            row = cache.get(ticker, {})

            if row.get('Status') != 'complete':
                writer.writerow({
                    col: row.get(col, '') if col != 'Ticker' else ticker
                    for col in CACHE_FIELDS
                })

    print(f'Added {changed} missing cells; output: {args.output}')

    print(
        f'Progress cache: {args.cache}; '
        f'unresolved report: {failed_path}'
    )

    print(
        'If rate-limited or interrupted, '
        're-run the same command to resume.'
    )


# ============================================================
# PROGRAM ENTRY POINT
# ============================================================
# Python executes the main() function when this script
# is run directly.
#
# This also allows the functions above to be imported
# into another Python program without automatically
# executing the entire Yahoo Finance lookup process.


if __name__ == '__main__':
    main()
